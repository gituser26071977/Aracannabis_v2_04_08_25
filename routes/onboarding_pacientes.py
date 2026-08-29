"""Onboarding de pacientes — cadastro administrativo com IA + fila de pendências.

Fluxo:
    POST /api/onboarding/paciente/sugerir  — extrai dados de texto livre (IA/heurística)
    POST /api/onboarding/paciente          — cadastra ou abre pendência (duplicado/incompleto)
    GET  /api/onboarding/pendentes         — fila para o administrativo
    POST /api/onboarding/pendentes/<id>/confirmar — cria ou usa existente
    POST /api/onboarding/pendentes/<id>/descartar — descarta
    POST /api/onboarding/documento/upload  — upload com IA (Gemini Vision) + pipeline completo
    POST /api/onboarding/documento/upload-batch — upload de planilha/lista em lote
    GET  /api/onboarding/documentos        — lista documentos do onboarding
"""

from __future__ import annotations

import logging
import os

from flask import Blueprint, jsonify, request
from flask_jwt_extended import jwt_required, get_jwt_identity

from models import Profissional, OnboardingDocumento
from services.onboarding_pacientes import (
    confirmar_pendencia,
    descartar_pendencia,
    listar_pendentes,
    registrar_paciente,
    sugerir_dados,
)

logger = logging.getLogger(__name__)

onboarding_bp = Blueprint("onboarding_pacientes", __name__)


def _current_user():
    user_id = get_jwt_identity()
    return Profissional.query.get(int(user_id)) if str(user_id).isdigit() else None


@onboarding_bp.route("/paciente/sugerir", methods=["POST"])
@jwt_required()
def sugerir_paciente():
    """Extrai dados do paciente a partir de texto livre (IA com fallback heurístico)."""
    data = request.get_json(silent=True) or {}
    texto = (data.get("texto") or "").strip()
    if not texto:
        return jsonify({"error": "informe o texto para sugestão"}), 400
    try:
        sugestao = sugerir_dados(texto)
    except Exception:  # noqa: BLE001
        logger.exception("sugestao_falhou")
        return jsonify({"error": "falha ao sugerir dados"}), 500
    return jsonify({"sugestao": sugestao}), 200


@onboarding_bp.route("/paciente", methods=["POST"])
@jwt_required()
def cadastrar_paciente():
    """Cadastra paciente. Se duplicado ou incompleto, abre pendência."""
    user = _current_user()
    if not user:
        return jsonify({"error": "Usuário não encontrado"}), 404
    data = request.get_json(silent=True) or {}
    try:
        resultado = registrar_paciente(
            data,
            origem="admin",
            criado_por=str(user.id),
            documento_id=int(data["documento_id"]) if data.get("documento_id") else None,
        )
    except Exception:  # noqa: BLE001
        logger.exception("onboarding_paciente_falhou")
        return jsonify({"error": "erro ao registrar paciente"}), 500

    if resultado["status"] == "criado":
        return jsonify({"message": "Paciente cadastrado", "resultado": resultado}), 201
    return jsonify({"message": "Pendência criada para revisão", "resultado": resultado}), 200


@onboarding_bp.route("/pendentes", methods=["GET"])
@jwt_required()
def listar_pendencias():
    """Fila de pendências do onboarding de pacientes."""
    try:
        itens = listar_pendentes()
    except Exception:  # noqa: BLE001
        return jsonify({"error": "erro ao listar pendências"}), 500
    return jsonify({"total": len(itens), "pendentes": [i.to_dict() for i in itens]}), 200


@onboarding_bp.route("/pendentes/<int:onboarding_id>/confirmar", methods=["POST"])
@jwt_required()
def confirmar(onboarding_id):
    user = _current_user()
    if not user:
        return jsonify({"error": "Usuário não encontrado"}), 404
    data = request.get_json(silent=True) or {}
    acao = data.get("acao", "criar")
    try:
        resultado = confirmar_pendencia(
            onboarding_id,
            acao=acao,
            dados=data.get("dados"),
            criado_por=str(user.id),
        )
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception:  # noqa: BLE001
        logger.exception("confirmar_pendencia_falhou")
        return jsonify({"error": "erro ao confirmar pendência"}), 500
    return jsonify({"message": "Pendência resolvida", "resultado": resultado}), 200


@onboarding_bp.route("/pendentes/<int:onboarding_id>/descartar", methods=["POST"])
@jwt_required()
def descartar(onboarding_id):
    user = _current_user()
    if not user:
        return jsonify({"error": "Usuário não encontrado"}), 404
    try:
        descartar_pendencia(onboarding_id, criado_por=str(user.id))
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception:  # noqa: BLE001
        logger.exception("descartar_pendencia_falhou")
        return jsonify({"error": "erro ao descartar pendência"}), 500
    return jsonify({"message": "Pendência descartada"}), 200


@onboarding_bp.route("/paciente/upload", methods=["POST"])
@jwt_required()
def upload_documento():
    """Upload de documento (imagem/PDF) no onboarding → OCR → sugestão.

    Salva o arquivo em uploads/onboarding/ e cria o registro
    `onboarding_documentos` (paciente_id NULL até cadastrar).
    """
    user = _current_user()
    if not user:
        return jsonify({"error": "Usuário não encontrado"}), 404

    from services.onboarding_pacientes import (
        ALLOWED_EXTENSIONS,
        _extensao,
        salvar_documento,
        sugerir_dados_de_documento,
    )

    arquivo = request.files.get("file")
    if not arquivo or not arquivo.filename:
        return jsonify({"error": "enviar o arquivo no campo 'file'"}), 400
    ext = _extensao(arquivo.filename)
    if ext not in ALLOWED_EXTENSIONS:
        return jsonify({"error": f"formato inválido. Permitidos: {sorted(ALLOWED_EXTENSIONS)}"}), 400

    conteudo = arquivo.read()
    if not conteudo:
        return jsonify({"error": "arquivo vazio"}), 400

    try:
        resultado = sugerir_dados_de_documento(conteudo, arquivo.filename)
    except Exception:  # noqa: BLE001
        logger.exception("upload_documento_ocr_falhou")
        return jsonify({"error": "erro ao processar o documento"}), 500

    try:
        doc = salvar_documento(
            conteudo,
            arquivo.filename,
            mime=arquivo.mimetype or "application/octet-stream",
            texto_extraido=resultado["texto_extraido"],
            confianca=resultado["confianca"],
            criado_por=str(user.id),
        )
    except Exception:  # noqa: BLE001
        logger.exception("upload_documento_salvar_falhou")
        return jsonify({"error": "erro ao salvar o documento"}), 500

    return jsonify({
        "sugestao": resultado["sugestao"],
        "texto_extraido": resultado["texto_extraido"],
        "confianca": resultado["confianca"],
        "documento_id": doc.id,
    }), 200


@onboarding_bp.route("/documento/upload", methods=["POST"])
@jwt_required()
def upload_documento_inteligente():
    """Upload de documento com analise inteligente (Gemini Vision) + pipeline completo.

    Aceita imagem (base64 via JSON ou multipart) e retorna:
    - tipo do documento
    - dados extraidos do paciente
    - se criou paciente, abriu pendencia ou detectou duplicado
    """
    user = _current_user()
    if not user:
        return jsonify({"error": "Usuário não encontrado"}), 404

    from services.intelligent_onboarding_service import onboarding_service

    ALLOWED_EXTENSIONS = {"png", "jpg", "jpeg", "pdf"}

    arquivo = request.files.get("file")
    conteudo = None
    filename = None

    if arquivo and arquivo.filename:
        ext = arquivo.filename.rsplit(".", 1)[-1].lower() if "." in arquivo.filename else ""
        if ext not in ALLOWED_EXTENSIONS:
            return jsonify({"error": f"Formato invalido. Permitidos: {sorted(ALLOWED_EXTENSIONS)}"}), 400
        conteudo = arquivo.read()
        filename = arquivo.filename
    else:
        data = request.get_json(silent=True) or {}
        imagem_b64 = data.get("imagem_base64")
        if not imagem_b64:
            return jsonify({"error": "Enviar arquivo (campo 'file') ou imagem_base64 no JSON"}), 400
        import base64 as b64_mod
        conteudo = b64_mod.b64decode(imagem_b64)
        filename = data.get("filename", "camera_capture.jpg")

    if not conteudo:
        return jsonify({"error": "Arquivo vazio"}), 400

    try:
        resultado = onboarding_service.processar_documento(
            conteudo, filename, criado_por=str(user.id)
        )
    except Exception:
        logger.exception("upload_inteligente_falhou")
        return jsonify({"error": "Erro ao processar documento"}), 500

    return jsonify(resultado), 201 if resultado.get("status") == "criado" else 200


@onboarding_bp.route("/documento/upload-batch", methods=["POST"])
@jwt_required()
def upload_lote():
    """Upload de planilha/lista com multiplos pacientes (processamento em lote).

    Aceita CSV ou XLSX com colunas: nome, cpf, telefone, email (ou imagem/PDF com lista).
    """
    user = _current_user()
    if not user:
        return jsonify({"error": "Usuário não encontrado"}), 404

    from services.intelligent_onboarding_service import onboarding_service

    arquivo = request.files.get("file")
    if not arquivo or not arquivo.filename:
        return jsonify({"error": "Enviar o arquivo no campo 'file'"}), 400

    ext = arquivo.filename.rsplit(".", 1)[-1].lower() if "." in arquivo.filename else ""

    if ext in ("csv", "xlsx"):
        return _processar_planilha(arquivo, ext, str(user.id))
    return _processar_imagem_lote(arquivo, str(user.id))


def _processar_planilha(arquivo, ext: str, criado_por: str):
    """Processa CSV ou XLSX com lista de pacientes."""
    conteudo = arquivo.read()
    logger.info(f"Processando planilha: {arquivo.filename} ({len(conteudo)} bytes)")

    if ext == "csv":
        import csv
        import io
        reader = csv.DictReader(io.StringIO(conteudo.decode("utf-8", errors="replace")))
        pacientes = list(reader)
    else:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True)
        ws = wb.active
        headers = [cell.value for cell in next(ws.iter_rows())]
        pacientes = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            pacientes.append(dict(zip(headers, [str(v) if v is not None else "" for v in row])))

    from services.intelligent_onboarding_service import onboarding_service

    resultados = []
    criados = 0
    pendentes = 0

    for dados in pacientes:
        nome = dados.get("nome", "").strip()
        if not nome:
            continue
        try:
            duplicados = onboarding_service._detectar_duplicados(dados)
            if duplicados:
                pendentes += 1
                resultados.append({
                    "status": "pendente",
                    "dados": {"nome": nome, "cpf": dados.get("cpf")},
                    "motivo": "duplicado",
                    "paciente_existente": duplicados[0].nome,
                    "paciente_id_existente": duplicados[0].id,
                })
                continue

            paciente = onboarding_service._criar_paciente(dados)
            compliance = onboarding_service._verificar_compliance(paciente)
            criados += 1
            resultados.append({
                "status": "criado",
                "dados": {"nome": nome, "cpf": dados.get("cpf")},
                "paciente_id": paciente.id,
                "compliance": compliance,
            })
        except Exception as e:
            logger.error(f"Erro processando paciente {nome}: {e}")
            resultados.append({
                "status": "erro",
                "dados": {"nome": nome},
                "erro": str(e),
            })

    return jsonify({
        "status": "lote_processado",
        "total": len(pacientes),
        "criados": criados,
        "pendentes": pendentes,
        "erros": sum(1 for r in resultados if r["status"] == "erro"),
        "resultados": resultados,
    }), 200


def _processar_imagem_lote(arquivo, criado_por: str):
    """Processa imagem/PDF com lista de pacientes (delega ao serviço inteligente)."""
    from services.intelligent_onboarding_service import onboarding_service

    analise = onboarding_service.analisar_documento(arquivo.read(), arquivo.filename)

    if not analise.get("multiplos_pacientes") or not analise.get("pacientes"):
        return jsonify({
            "status": "nao_identificado_como_lista",
            "analise": analise,
            "mensagem": "Documento nao identificado como lista de pacientes. Use /documento/upload para upload individual.",
        }), 200

    resultado_lote = onboarding_service._processar_lote(
        analise, arquivo.read(), arquivo.filename, criado_por
    )
    return jsonify(resultado_lote), 200


@onboarding_bp.route("/documentos", methods=["GET"])
@jwt_required()
def listar_documentos():
    """Lista documentos enviados no onboarding."""
    docs = OnboardingDocumento.query.order_by(OnboardingDocumento.created_at.desc()).limit(100).all()
    return jsonify({
        "total": len(docs),
        "documentos": [d.to_dict() for d in docs],
    }), 200


@onboarding_bp.route("/documentos/<int:documento_id>/arquivo", methods=["GET"])
@jwt_required()
def baixar_documento(documento_id):
    """Download do documento enviado no onboarding."""
    from models import OnboardingDocumento
    from flask import send_file

    doc = OnboardingDocumento.query.get(documento_id)
    if not doc or not os.path.exists(doc.caminho_arquivo):
        return jsonify({"error": "documento não encontrado"}), 404
    return send_file(doc.caminho_arquivo, mimetype=doc.mime or "application/octet-stream", as_attachment=True, download_name=doc.nome_original)
